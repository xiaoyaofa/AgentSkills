#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""table_to_xlsx_openpyxl.py — 用 openpyxl 直接把 TSV 写成 .xlsx（【免 LibreOffice】）。

替代 table_to_xlsx.py 的"FODS + LibreOffice"路线；pip install openpyxl 即可，
Windows/Linux 一致，不再需要任何系统软件。

用法（与 table_to_xlsx.py 完全一致）:
    table_to_xlsx_openpyxl.py 输出.xlsx \
        --sheet "0-说明:notes.tsv:2.6,11" \
        --sheet "1-已用信号:used.tsv:2,3.2,2.2,1.6,2.8,4.2"
    （--sheet 语法: "表名:TSV文件[:列宽cm,用逗号分隔]"）
"""

# ---------------------------------------------------------------------------
# ★ 强制 UTF-8 输出（Windows 中文环境必需）
#   Windows 中文环境默认编码是 GBK(cp936)：脚本里只要有箭头/数学符号（`↔` `→` `≤`）就会
#   UnicodeEncodeError 直接崩 —— **而且是跑到打印那一步才崩，前面算完的结果一条都出不来**，
#   看着像"检查没通过"，实际是"检查没执行"。实测：`check_deliverable_consistency.py` 崩在
#   `print(report)`、`verify_pinmap.py` 崩在打印 `[A] 引脚↔信号名` 那行。
#   ★ 被 `-I` 拉起时（`-I` 蕴含 `-E`）PYTHONIOENCODING / PYTHONUTF8 都无效，只能在进程内改。
#   ★ **不要**改成 import 一个共享小模块：`-I` 蕴含 `-P`，脚本目录不进 sys.path，只会静默失效；
#     所以这段在每个脚本里内联一份（自带 import，放在 docstring 之后、其余 import 之前也能跑）。
# ---------------------------------------------------------------------------
import sys as _sys

for _s in (_sys.stdout, _sys.stderr):
    try:
        _s.reconfigure(encoding='utf-8', errors='replace')
    except (AttributeError, ValueError):        # 被重定向成非文本流时忽略
        pass
if _sys.platform == 'win32':
    try:                                        # 控制台代码页也切到 UTF-8，原生 cmd 窗口才不会显示成乱码
        import ctypes
        ctypes.windll.kernel32.SetConsoleOutputCP(65001)
    except Exception:
        pass


import argparse
import os
import re
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit('缺 openpyxl：请先  pip install openpyxl')


WIDTHS_RX = re.compile(r'^[\d.]+(?:,[\d.]+)*$')


def parse_sheet_spec(spec):
    """解析 --sheet 的 "表名:路径[:列宽]"。

    ★ 不能用朴素的 spec.split(':')：Windows 路径带盘符（D:/work/x.tsv）会被切碎，
      取到 'D' 当路径 → 报"找不到 D"。规则：
        第一个 ':' 之前是表名；
        末尾那段**只有**形如 "2,3.2" 才当列宽（且要求候选路径存在，或该段含 ',' / '.'）；
        剩下整段都是路径（含盘符）。

    返回 (name, path, widths)。与 table_to_xlsx.py 里同名函数保持一致。
    """
    if ':' not in spec:
        sys.exit(f'--sheet 语法错误（应为 "表名:TSV文件[:列宽]"）: {spec}')
    name, rest = spec.split(':', 1)
    widths = []
    if ':' in rest:
        cand_path, cand_w = rest.rsplit(':', 1)
        looks_like_path = ('/' in cand_path or '\\' in cand_path or '.' in cand_path)
        if WIDTHS_RX.match(cand_w) and (os.path.isfile(cand_path) or ',' in cand_w
                                        or '.' in cand_w or looks_like_path):
            rest, widths = cand_path, [float(x) for x in cand_w.split(',')]
    if not name:
        sys.exit(f'--sheet 表名为空: {spec}')
    return name, rest, widths


def read_tsv(path):
    rows = []
    with open(path, encoding='utf-8', errors='replace') as f:
        for line in f:
            rows.append(line.rstrip('\r\n').split('\t'))
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('out')
    ap.add_argument('--sheet', action='append', required=True,
                    metavar='NAME:FILE[:W1,W2,...]')
    a = ap.parse_args()

    wb = Workbook()
    wb.remove(wb.active)                       # 去掉默认空白 sheet

    head_font = Font(bold=True)
    head_fill = PatternFill('solid', fgColor='D9E2F3')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for si, spec in enumerate(a.sheet):
        name, path, widths = parse_sheet_spec(spec)
        if not os.path.isfile(path):
            sys.exit(f'找不到文件 {path}（--sheet 给的是 "{spec}"）')
        rows = read_tsv(path)
        ws = wb.create_sheet(title=name[:31] or f'S{si}')
        for r in rows:
            ws.append(r)
        for j in range(1, (max(len(r) for r in rows) if rows else 1) + 1):
            c = ws.cell(row=1, column=j)
            c.font = head_font
            c.fill = head_fill
            c.alignment = Alignment(vertical='center')
        for row in ws.iter_rows(min_row=1, max_row=len(rows),
                                max_col=ws.max_column):
            for c in row:
                c.border = border
        for j, w in enumerate(widths, start=1):    # cm -> 近似字符宽
            ws.column_dimensions[get_column_letter(j)].width = round(w * 5.4, 1)
        ws.freeze_panes = 'A2'
        print(f'   {name}: {max(len(rows) - 1, 0)} 行数据')

    wb.save(a.out)
    print(f'XLSX: {a.out}  ({len(a.sheet)} sheet, openpyxl 直写)')


if __name__ == '__main__':
    main()
